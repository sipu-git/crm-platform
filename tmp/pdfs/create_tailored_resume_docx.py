from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


OUT = Path(r"C:\CRM Platform\output\docx\Sipu_Rana_NIQ_Associate_Software_Engineer_Resume.docx")

NAVY = "202C3D"
BODY = "4D5A6D"
LIGHT_BLUE = "C7D3E1"
RULE = "252F3E"


def set_font(run, name="Calibri", size=10, color=BODY, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def set_bottom_border(paragraph, color=RULE, size="6", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure(p, before=0, after=0, line=12.3, align=None, keep_next=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)
    fmt.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    return p


def add_text_paragraph(doc, text, size=10.05, line=12.45, before=0, after=0, align=None):
    p = doc.add_paragraph()
    configure(p, before, after, line, align)
    set_font(p.add_run(text), size=size)
    return p


def add_section(doc, title, before=2.2, after=4.1):
    p = doc.add_paragraph()
    configure(p, before=before, after=after, line=11.5, keep_next=True)
    set_font(p.add_run(title.upper()), size=10.9, color=NAVY, bold=True)
    set_bottom_border(p)
    return p


def add_skill(doc, label, text):
    p = doc.add_paragraph()
    configure(p, before=0, after=0, line=12.35)
    set_font(p.add_run(label), size=9.82, color=NAVY, bold=True)
    set_font(p.add_run(" " + text), size=9.82)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    configure(p, before=0, after=0, line=12.25)
    p.paragraph_format.left_indent = Mm(12.5)
    p.paragraph_format.first_line_indent = Mm(-5.0)
    set_font(p.add_run("•"), size=11.5, color="000000", bold=True)
    set_font(p.add_run("  " + text), size=10.0)
    return p


def add_project_title(doc, title, before=0.8, after=0.5):
    p = doc.add_paragraph()
    configure(p, before=before, after=after, line=11.8, keep_next=True)
    set_font(p.add_run(title), size=10.55, color=NAVY, bold=True)
    return p


def add_education_line(doc, label, org):
    p = doc.add_paragraph()
    configure(p, before=0, after=0, line=11.75)
    set_font(p.add_run(label), size=10.55, color=NAVY, bold=True)
    set_font(p.add_run(" -  " + org), size=10.45)
    return p


def add_education_detail(doc, detail):
    p = doc.add_paragraph()
    configure(p, before=0, after=0.9, line=11.5)
    set_font(p.add_run(detail), size=9.9, italic=True)
    return p


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(11.3)
    section.right_margin = Mm(11.3)
    section.top_margin = Mm(15.5)
    section.bottom_margin = Mm(11.0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    doc.core_properties.title = "Sipu Rana Associate Software Engineer Resume"
    doc.core_properties.author = "Sipu Rana"
    doc.core_properties.subject = "Tailored resume for Associate Software Engineer roles"

    # Header
    p = doc.add_paragraph()
    configure(p, before=0, after=3.6, line=20.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("SIPU RANA"), size=17.3, color=NAVY, bold=True)

    p = doc.add_paragraph()
    configure(p, before=0, after=9.8, line=12.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    contact = "+91 9827505917  |  sipurana38@gmail.com  |  https://www.github.com/sipu-git  |  https://www.linkedin.com/in/sipu-rana-72b56b242"
    set_font(p.add_run(contact), size=9.3)
    set_bottom_border(p, color=LIGHT_BLUE, size="7", space="6")

    add_section(doc, "Career Summary", before=0, after=4.2)
    add_text_paragraph(
        doc,
        "Associate software engineer with 1+ year of experience building, debugging, and optimizing responsive, component-based web applications across the full stack (JavaScript/TypeScript, React.js, Next.js, Node.js, Express.js). Familiar with Angular fundamentals and modern web development practices. Skilled in REST API integration, state handling, Git/GitHub workflows, GitHub Actions CI/CD, code reviews, and Agile/Scrum collaboration; focused on clean, maintainable, secure, high-performance software.",
        size=10.15,
        line=12.65,
        after=2.0,
    )

    add_section(doc, "Technical Skills", before=0.5, after=3.2)
    skills = [
        ("Version Control & Collaboration:", "Git, GitHub (branching, pull requests, code review, merge conflict resolution), GitHub Actions, CI/CD pipelines, Agile/Scrum"),
        ("Testing & Code Quality:", "Jest, Karma/Jasmine (familiarity), Postman API validation, error handling, cross-layer debugging (frontend, backend, database)"),
        ("Languages & Frontend:", "Angular (component-based architecture), JavaScript (ES6+), TypeScript, HTML5, CSS3, SCSS, React.js, Next.js (SSR/SSG, App Router), TanStack Query (React Query), Redux Toolkit"),
        ("Backend & APIs:", "Node.js, Express.js, RESTful API design & integration, JWT/OAuth, middleware"),
        ("Databases:", "MongoDB & Mongoose, PostgreSQL, MySQL, Prisma ORM, schema design, Redis"),
        ("Cloud & DevOps:", "AWS (EC2, S3, RDS, IAM, CloudWatch), Docker, CI/CD"),
    ]
    for label, value in skills:
        add_skill(doc, label, value)

    add_section(doc, "Professional Experience", before=1.6, after=3.8)
    p = doc.add_paragraph()
    configure(p, before=0, after=2.4, line=12.8)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(section.page_width - section.left_margin - section.right_margin, WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run("Software Developer  -  Sparkradix Technologies Pvt. Ltd"), size=10.8, color=NAVY, bold=True)
    set_font(p.add_run("\t2025 - Present"), size=10.45)
    add_text_paragraph(
        doc,
        "Contribute to production React.js/Next.js applications and backend services; collaborate in Agile/Scrum teams, open and review pull requests, and resolve merge conflicts to keep releases reliable.",
        size=10.05,
        line=12.45,
        after=0.7,
    )
    for item in [
        "Debug and optimize frontend, backend, and database code; validate REST APIs with Postman and resolve defects before release.",
        "Design and implement RESTful APIs with Express.js, integrating relational (MySQL, PostgreSQL via Prisma) and NoSQL (MongoDB) data layers.",
        "Build reusable, responsive, cross-browser UI components and dynamic templates using semantic HTML5, CSS3/SCSS, JavaScript, and TypeScript.",
        "Collaborate with cross-functional teams to translate requirements into shippable features and incorporate code-review feedback to improve maintainability, security, and performance.",
    ]:
        add_bullet(doc, item)

    add_section(doc, "Projects", before=1.3, after=3.8)
    add_project_title(doc, "Intelligent Customer Relationship & Workflow Management System", before=0, after=0.4)
    add_bullet(doc, "Built an AI-powered CRM and business automation platform with React.js, TypeScript, Node.js, Express.js, Prisma, PostgreSQL/AWS RDS, Redis, Tailwind CSS, and role-based access control for lead tracking and sales workflows.")
    add_bullet(doc, "Integrated TanStack Query (React Query) for backend REST API fetching, caching, mutations, invalidation, and loading/error state handling; used Redux Toolkit for client-side UI state and integrated WhatsApp Business API plus Gmail/Zoho Mail.")
    add_project_title(doc, "Expense Tracking Web Application  &  Role-Based Access HRMS Web Application", before=0.9, after=0.4)
    add_bullet(doc, "Built an expense platform with React.js, Redux Toolkit, Express.js, PostgreSQL, AI-powered auto-categorization, and Redis-based OTP authentication.")
    add_bullet(doc, "Built a scalable HRMS platform with RBAC for employee, attendance, leave, and payroll management, optimized with Redis caching and Docker.")

    add_section(doc, "Education", before=1.4, after=3.2)
    add_education_line(doc, "Master in Computer Application", "Gandhi Engineering College")
    add_education_detail(doc, "CGPA: 8.38  |  2023 - 2025")
    add_education_line(doc, "Bachelor in Computer Application", "BCCM College, Berhampur University")
    add_education_detail(doc, "Score: 79%  |  2020 - 2023")

    add_section(doc, "Courses & Certifications", before=1.1, after=3.0)
    p = doc.add_paragraph()
    configure(p, before=0, after=0.4, line=11.95)
    set_font(p.add_run("Full Stack Software Development (MERN Stack)"), size=10.35, color=NAVY, bold=True)
    set_font(p.add_run("  -  Tetratrion Technologies Pvt. Ltd. (Sept 2024 - Apr 2025)"), size=10.2)
    p = doc.add_paragraph()
    configure(p, before=0, after=0, line=11.4)
    set_font(p.add_run("Full Stack Development · RESTful API & Backend Architecture · System & Database Design · Frontend Testing & Collaboration"), size=9.65, italic=True)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
